"""Karussell: projekte/<name>/inhalt.json -> Bilder fuer Instagram ODER eine Datei fuer Canva.

  python bauen.py                   aktuelles Projekt komplett bauen
  python bauen.py --projekt kompass ein bestimmtes Projekt bauen
  python bauen.py --cockpit         Cockpit oeffnen (aendern, speichern, rendern)
  python bauen.py --pdf [--projekt name]   nur das PDF bauen
  python bauen.py --vorschau        nur anschauen, ohne Regler

Ergebnis je Projekt in ../export/<name>/ (bilder/, karussell.pptx, kontaktbogen.png).
"""
import json, re, sys, os, io, time, zlib, shutil, zipfile, subprocess, socketserver, threading, webbrowser, functools
import base64, hashlib
import http.server
from pathlib import Path
from urllib.parse import urlsplit, parse_qs
from playwright.sync_api import sync_playwright
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont, ImageOps

import projekt as P

HIER     = P.HIER
PX       = 9525                   # EMU pro Pixel
SPERRE   = threading.Lock()
STIL_STD = P.STIL_STD             # 4:5. Alte Projekte tragen ihr 3:4 im eigenen stil.
BILDARTEN = {".png", ".jpg", ".jpeg", ".webp", ".gif", ".svg"}   # was /bild annimmt
WANDELN = {".heic", ".heif", ".tif", ".tiff"}                     # iPhone-Fotos: auf dem Mac per sips zu JPG


def _mac_wandeln(daten, datei):
    """HEIC/TIFF mit dem macOS-eigenen sips in JPG wandeln (kein Zusatzpaket).
    EXIF und GPS entfernt danach _aufbereiten wie bei jedem Bild."""
    import tempfile
    with tempfile.TemporaryDirectory() as t:
        q, z = Path(t) / ("q" + Path(datei).suffix.lower()), Path(t) / "z.jpg"
        q.write_bytes(daten)
        r = subprocess.run(["sips", "-s", "format", "jpeg", str(q), "--out", str(z)],
                           capture_output=True, timeout=60)
        if r.returncode or not z.is_file():
            raise ValueError("Bild liess sich nicht umwandeln")
        return z.read_bytes(), Path(datei).stem + ".jpg"


def browser_auf(url):
    """Auf dem Mac lieber Chrome (rechnet Umbrueche wie der Export), sonst der Standardbrowser."""
    if sys.platform == "darwin":
        try:
            if subprocess.run(["open", "-a", "Google Chrome", url], capture_output=True,
                              timeout=20).returncode == 0:
                return
        except Exception:
            pass
    webbrowser.open(url)

# Auszeichnungen als Tags, beliebig ueberlagerbar. Muss zu render.js passen.
TAG = re.compile(r"<(/?)(b|i|u|m|c|f)(?:=([^>]{0,64}))?>")


def laden(name=None):
    d = P.lesen(name)
    return d["slides"], P.stil(d)


def segmente(s):
    """-> [(text, format)] mit format aus b/i/u/m/c/f"""
    s = s or ""
    out, st, start = [], {}, 0
    for m in TAG.finditer(s):                 # linear, auch bei "<f=<f=<f=..." ohne ">"
        if m.start() > start:
            out.append((s[start:m.start()], dict(st)))
        if m.group(1):
            st.pop(m.group(2), None)
        else:
            st[m.group(2)] = True if m.group(3) is None else m.group(3)
        start = m.end()
    if start < len(s):
        out.append((s[start:], dict(st)))
    return out or [("", {})]


# ---------------------------------------------------------------- Server

def stand():
    """'<aktuelles Projekt>|<Pruefsumme seiner inhalt.json>'. Das Cockpit
    erkennt daran Aenderungen von aussen und zieht selbst nach, auch wenn
    von aussen ein anderes Projekt geoeffnet wurde."""
    return stand_von(P.aktuelles())


def stand_von(name):
    # Pruefsumme statt Zeitstempel: D: ist exFAT, dort springt die Uhrzeit nur
    # in ganzen Sekunden, zwei Aenderungen in derselben Sekunde gingen unter.
    try:
        return "%s|%08x" % (name, zlib.crc32((P.ordner(name) / "inhalt.json").read_bytes()))
    except (OSError, ValueError):
        return "%s|0" % (name or "")


SCHRIFT_SPERRE = threading.Lock()


def vorlagen_json():
    aus = []
    for n in P.vorlagen():
        try:
            o = P.vorlage_ordner(n)
            d = json.loads((o / "inhalt.json").read_text(encoding="utf-8-sig"))
            st = P.stil(d)
            eigen = o.parent == P.VORLAGEN
            titel, text = _vorlagen_info(d, eigen)
            aus.append({"name": n, "slides": len(d.get("slides") or []), "breite": int(st["breite"]),
                        "hoehe": int(st["hoehe"]), "vorschau": (o / P.VORSCHAU).exists(),
                        "folien": (o / P.FOLIEN).exists(), "titel": titel, "beschreibung": text,
                        "eigen": eigen})
        except Exception:
            continue
    return json.dumps(aus, ensure_ascii=False)


def _vorlagen_info(d, eigen):
    """Titel und Beschreibung aus inhalt.json ("vorlage"), nur kurzer sichtbarer Text.
    Eigene Vorlagen haben keine (als_vorlage nimmt sie heraus); traegt eine
    trotzdem welche, gibt sie sich als Kit-Vorlage aus und bleibt beim Namen."""
    i = d.get("vorlage") if isinstance(d, dict) else None
    if eigen or not isinstance(i, dict):
        return "", ""
    return P.kurztext(i.get("titel"), 40), P.kurztext(i.get("beschreibung"), 140)


def vorschau_bauen(port, vorlage, ordner=None):
    """Folie 1 der Vorlage als kleines Bild (270 px breit) fuer die Galerie, dazu alle
    Folien nebeneinander (.folien.png). Leere Bildrahmen erscheinen als ruhige Flaeche.
    ordner: nur fuer die mitgelieferten Vorlagen im Kit, sonst die eigenen."""
    o = ordner or (P.VORLAGEN / vorlage)
    if ordner and P.vorlage_ordner(vorlage) != ordner:
        # vorlage.html laedt dann die gleichnamige eigene Vorlage, und deren Bilder
        # landeten in den Vorschaubildern des Kits, die an alle gehen
        raise ValueError("Eine eigene Vorlage gleichen Namens verdeckt die Kit-Vorlage %s" % vorlage)
    d = json.loads((o / "inhalt.json").read_text(encoding="utf-8-sig"))
    B, H = format_ok(P.stil(d))
    n = min(len(d.get("slides") or []), 20)
    with sync_playwright() as pw:
        b = _browser(pw)
        pg = b.new_page(viewport={"width": B, "height": H}, device_scale_factor=270 / B)
        pg.goto(f"http://127.0.0.1:{port}/vorlage.html?vorlage={vorlage}&render=0&platzhalter=1")
        pg.wait_for_function(FERTIG, timeout=20000)
        pg.screenshot(path=str(o / P.VORSCHAU))
        pg = b.new_page(viewport={"width": max(300, 20 + n * 226), "height": 300})
        pg.goto(f"http://127.0.0.1:{port}/vorlage.html?vorlage={vorlage}&platzhalter=1&streifen=1")
        pg.wait_for_function(FERTIG, timeout=30000)
        pg.screenshot(path=str(o / P.FOLIEN), full_page=True)
        b.close()


FERTIG = "() => window.FERTIG === true"      # Funktionsform: die Seiten-CSP verbietet eval

FORMAT_ZIEL = {"3:4": 1440, "34": 1440, "1440": 1440, "4:5": 1350, "45": 1350, "1350": 1350}


def format_umstellen(port, name, ziel):
    """Karussell auf 3:4 oder 4:5 umstellen (k.py format). Gerechnet wird mit
    formatUmstellen aus render.js in vorlage.html, mit den echten Schriften,
    damit k.py genau dasselbe tut wie die Formatwahl im Cockpit."""
    hoehe = FORMAT_ZIEL.get(str(ziel).strip())
    if not hoehe:
        raise ValueError("Format bitte als 3:4 oder 4:5")
    ordner = P.ordner(name)
    ziel_datei = ordner / "inhalt.json"
    if ziel_datei.is_symlink() or P._verknuepft(ordner):     # wie beim Speichern: nie durch Links schreiben
        raise ValueError("inhalt.json oder der Projektordner ist eine Verknuepfung")
    vorher = stand_von(name)
    d = P.lesen(name)
    B, H = format_ok(P.stil(d))
    wort = "3:4" if hoehe == 1440 else "4:5"
    if B != 1080:
        raise ValueError("Umstellen geht nur bei Folien mit 1080 px Breite")
    if len(d.get("slides") or []) > 100:
        raise ValueError("Zu viele Folien (hoechstens 100)")
    if H == hoehe:
        return "Ist schon %s (%d x %d)." % (wort, B, H)
    with sync_playwright() as pw:
        b = _browser(pw)
        try:
            pg = b.new_page(viewport={"width": B, "height": H})
            pg.goto(f"http://127.0.0.1:{port}/vorlage.html?projekt={name}&format={hoehe}")
            pg.wait_for_function(FERTIG, timeout=60000)
            e = pg.evaluate("() => window.ERGEBNIS")
        finally:
            b.close()
    if not isinstance(e, dict) or e.get("fehler"):
        raise ValueError("Umstellen ging nicht: %s" % " ".join(str((e or {}).get("fehler", "keine Antwort")).split())[:200])
    neu, hinweise = e.get("daten"), e.get("hinweise")
    if not isinstance(neu, dict) or not isinstance(neu.get("slides"), list) \
            or len(neu["slides"]) != len(d.get("slides") or []) or P.stil(neu).get("hoehe") != hoehe:
        raise ValueError("Umstellen ging nicht: unerwartetes Ergebnis")
    if stand_von(name) != vorher:          # inzwischen im Cockpit gesichert: nichts ueberschreiben
        raise ValueError("Das Karussell wurde gerade geaendert. Bitte noch einmal umstellen.")
    P.schreiben(neu, name)
    zeilen = ["Umgestellt auf %s (%d x %d)." % (wort, B, hoehe)]
    hinweise = [h for h in (hinweise if isinstance(hinweise, list) else []) if isinstance(h, str)][:40]
    # je Hinweis genau eine Zeile: Umbrueche aus Projektdaten duerfen keine eigenen Zeilen bauen
    zeilen += ["Bitte ansehen:"] + ["  " + " ".join(h.split())[:200] for h in hinweise] if hinweise else ["Alles passt."]
    return "\n".join(zeilen)


def format_ok(st):
    """Breite und Hoehe aus inhalt.json, nur in vernuenftigen Grenzen
    (fremde Dateien koennten sonst viele GB Speicher belegen)."""
    B, H = int(zahl(st.get("breite"), 0)), int(zahl(st.get("hoehe"), 0))
    if not (100 <= B <= 5000 and 100 <= H <= 8000):
        raise ValueError("Format ungueltig: %r x %r" % (str(st.get("breite"))[:20], str(st.get("hoehe"))[:20]))
    return B, H


def _browser(pw):
    """Chromium fuer Export und Vorschau, mit eigener Sandbox (Playwright
    schaltet sie sonst ab). Faellt auf ohne zurueck, falls das System sie
    nicht erlaubt (manche Linux-Container)."""
    try:
        return pw.chromium.launch(chromium_sandbox=True)
    except Exception as e:
        print("Hinweis: Browser ohne eigene Sandbox gestartet (%s)" % str(e).splitlines()[0][:120])
        return pw.chromium.launch()


def oeffnen(pfad, markieren=None):
    """Ordner im Explorer bzw. Finder zeigen, optional eine Datei darin markieren."""
    import platform
    system = platform.system()
    if system == "Windows":
        if markieren:
            subprocess.Popen(["explorer", "/select,", str(markieren)])
        else:
            os.startfile(str(pfad))
    elif system == "Darwin":
        subprocess.Popen(["open", "-R", str(markieren)] if markieren else ["open", str(pfad)])
    else:
        subprocess.Popen(["xdg-open", str(pfad)])


# Icons: im Kit als EINE Datei bib/icons.json ({"lucide/heart": "<svg ...>"}),
# statt 8.000 Einzeldateien. Gleiche Adressen wie die Einzeldateien.
ICON_WEG = re.compile(r"bib/((?:lucide|tabler|tabler-voll)/[a-z0-9-]{1,80})\.svg")
_ICONS = []


def icon_paket():
    if not _ICONS:
        try:
            _ICONS.append(json.loads((HIER / "bib" / "icons.json").read_text(encoding="utf-8")))
        except (OSError, ValueError):
            _ICONS.append({})
    return _ICONS[0]


LADEN = {"modell": None, "gelesen": 0, "gesamt": 0, "fehler": "", "fertig": ""}
LADE_SPERRE = threading.Lock()


def modelle_stand():
    import freisteller as F
    return {"da": F.vorhanden(), "mb": {m: q["mb"] for m, q in F.QUELLEN.items()},
            "laden": dict(LADEN)}


def modell_laden_starten(modell):
    """Freisteller im Hintergrund laden. Fortschritt ueber GET /modelle."""
    import freisteller as F
    import einrichten as E
    if modell in F.vorhanden():
        return {"ok": True, "da": True}
    if not LADE_SPERRE.acquire(blocking=False):
        return {"ok": False, "fehler": "es laedt schon ein Freisteller"}
    LADEN.update(modell=modell, gelesen=0, gesamt=F.QUELLEN[modell]["mb"] * 1_000_000,
                 fehler="", fertig="")

    def los():
        try:
            E.modell_laden(modell, lambda n, g: LADEN.update(gelesen=n, gesamt=g))
            LADEN["fertig"] = modell
        except Exception as e:
            LADEN["fehler"] = str(e)[:200]
        finally:
            LADEN["modell"] = None
            LADE_SPERRE.release()
    threading.Thread(target=los, daemon=True).start()
    return {"ok": True}


# ---------------------------------------------------------------- PDF
# Jede Folie als Bild in doppelter Aufloesung (sieht genau aus wie das
# Instagram-Bild, nur schaerfer), JPEG Qualitaet 95 ohne Farbunterabtastung.
# Die JPEGs gehen unveraendert ins PDF, ohne zweites Komprimieren.

def pdf_schreiben(ziel, jpegs, B, H):
    """Minimales PDF: eine Seite je Folie, Seitengroesse in Punkt = Pixel der Folie."""
    offs = []
    puffer = bytearray(b"%PDF-1.4\n%\xe2\xe3\xcf\xd3\n")

    def obj(nr, inhalt):
        offs.append((nr, len(puffer)))
        puffer.extend(b"%d 0 obj\n" % nr + inhalt + b"\nendobj\n")
    n = len(jpegs)
    seiten = [3 + 3 * i for i in range(n)]
    obj(1, b"<< /Type /Catalog /Pages 2 0 R >>")
    obj(2, b"<< /Type /Pages /Count %d /Kids [%s] >>" % (n, b" ".join(b"%d 0 R" % k for k in seiten)))
    for i, (jpg, (bw, bh)) in enumerate(jpegs):
        sn = seiten[i]
        inhalt = b"q %d 0 0 %d 0 0 cm /Im0 Do Q" % (B, H)
        obj(sn, b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 %d %d] /Resources << /XObject << /Im0 %d 0 R >> >> /Contents %d 0 R >>"
            % (B, H, sn + 1, sn + 2))
        obj(sn + 1, b"<< /Type /XObject /Subtype /Image /Width %d /Height %d /ColorSpace /DeviceRGB "
                    b"/BitsPerComponent 8 /Filter /DCTDecode /Length %d >>\nstream\n" % (bw, bh, len(jpg))
            + jpg + b"\nendstream")
        obj(sn + 2, b"<< /Length %d >>\nstream\n" % len(inhalt) + inhalt + b"\nendstream")
    xref = len(puffer)
    gesamt = 3 + 3 * n
    puffer.extend(b"xref\n0 %d\n0000000000 65535 f \n" % gesamt)
    for nr, off in sorted(offs):
        puffer.extend(b"%010d 00000 n \n" % off)
    puffer.extend(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (gesamt, xref))
    tmp = ziel.with_name(ziel.name + ".tmp")
    tmp.write_bytes(bytes(puffer))
    os.replace(tmp, ziel)


def pdf_bauen(port, name=None):
    name = name or P.aktuelles()
    slides, stil = laden(name)
    B, H = format_ok(stil)
    aus = P.export(name)
    aus.mkdir(parents=True, exist_ok=True)
    jpegs = []
    with sync_playwright() as pw:
        b = _browser(pw)
        pg = b.new_page(viewport={"width": B, "height": H}, device_scale_factor=2)
        for i in range(len(slides)):
            pg.goto(f"http://127.0.0.1:{port}/vorlage.html?projekt={name}&render={i}")
            pg.wait_for_function(FERTIG, timeout=20000)
            bild = Image.open(io.BytesIO(pg.screenshot())).convert("RGB")
            puffer = io.BytesIO()
            bild.save(puffer, "JPEG", quality=95, subsampling=0, optimize=True)
            jpegs.append((puffer.getvalue(), bild.size))
        b.close()
    ziel = aus / "karussell.pdf"
    pdf_schreiben(ziel, jpegs, B, H)
    return "PDF fertig (%d Seiten, %.1f MB)" % (len(jpegs), ziel.stat().st_size / 1e6)


def projekte_json():
    aus = []
    for n in P.liste():
        try:
            d = P.lesen(n)
            st = P.stil(d)
            aus.append({"name": n, "slides": len(d.get("slides") or []),
                        "breite": int(st["breite"]), "hoehe": int(st["hoehe"]),
                        "geaendert": int((P.ordner(n) / "inhalt.json").stat().st_mtime),
                        "groesse": (P.ordner(n) / "inhalt.json").stat().st_size})
        except Exception:
            aus.append({"name": n, "slides": 0, "breite": 0, "hoehe": 0})
    return json.dumps({"aktuell": P.aktuelles(), "projekte": aus}, ensure_ascii=False)


GROESSTE_KANTE = 3200       # reicht fuer 1080 breite Folien auch bei 3-fach Zoom


FORMAT = {".png": "PNG", ".jpg": "JPEG", ".jpeg": "JPEG", ".webp": "WEBP", ".gif": "GIF"}
MAX_PIXEL = 60_000_000      # Breite x Hoehe x Einzelbilder, darueber droht Speichernot (8 GB)
MAX_UPLOAD = 50_000_000     # Bytes je Anfrage


def _aufbereiten(daten, endung):
    """Prueft und bereitet ein hochgeladenes Rasterbild auf. Das Format wird an
    der Endung festgemacht, nicht am Inhalt (sonst kaemen EPS, EMF und Co.
    durch). Fotos mit Metadaten (GPS!) oder ueber GROESSTE_KANTE werden neu
    gespeichert: richtig gedreht, verkleinert, ohne EXIF. Alles andere bleibt
    bitgleich, damit keine Qualitaet verloren geht."""
    fmt = FORMAT[endung]
    with Image.open(io.BytesIO(daten), formats=[fmt]) as im:
        if im.width * im.height * getattr(im, "n_frames", 1) > MAX_PIXEL:
            raise ValueError("Bild zu gross (zu viele Pixel)")
        if fmt == "GIF":
            im.verify()
            return daten
        im.load()
        hat_meta = bool(im.getexif()) or "exif" in im.info or "xmp" in im.info
        if not hat_meta and max(im.size) <= GROESSTE_KANTE:
            return daten
        im = ImageOps.exif_transpose(im)
        if max(im.size) > GROESSTE_KANTE:
            f = GROESSTE_KANTE / max(im.size)
            im = im.resize((round(im.width * f), round(im.height * f)), Image.LANCZOS)
        puffer = io.BytesIO()
        if fmt == "JPEG":
            im.convert("RGB").save(puffer, "JPEG", quality=94, optimize=True)
        elif fmt == "WEBP":
            im.save(puffer, "WEBP", quality=94)
        else:
            im.save(puffer, "PNG")
        return puffer.getvalue()


def bild_ablegen(ordner, datei, daten):
    """Upload in den Projektordner. Gleicher Name, anderer Inhalt: neuer Name
    statt Ueberschreiben (sonst wechselt ein anderes Bild still sein Motiv)."""
    pfad = ordner / datei
    stamm, endung = pfad.stem, pfad.suffix
    e = endung.lower()
    # Nur echte Bilder: sonst steht spaeter ein leerer Kasten auf der Folie
    if e == ".svg":
        if "<svg" not in daten[:2048].decode("utf-8", "ignore").lower():
            raise ValueError("keine SVG-Grafik")
        inhalt = daten
    else:
        try:
            inhalt = _aufbereiten(daten, e)
        except ValueError:
            raise
        except Exception:
            raise ValueError("keine lesbare Bilddatei")
    n = 2
    while pfad.exists() and pfad.read_bytes() != inhalt:
        pfad = ordner / ("%s_%d%s" % (stamm, n, endung)); n += 1
    if pfad.is_symlink():
        raise ValueError("Ziel ist eine Verknuepfung")
    if not pfad.exists():
        pfad.write_bytes(inhalt)
    return pfad.name


GENAU_SPERRE = threading.Lock()      # der genaue Freisteller braucht ~3 GB, nie zwei zugleich
SCHNELL_SPERRE = threading.Lock()    # eine Grafikkarten-Sitzung vertraegt keine zwei Laeufe zugleich


def freistellen_auftrag(name, koerper):
    """{"datei": "...", "modell": "schnell|genau"} -> {"ok", "datei"} oder {"ok": False, "fehler"}.
    Schnell rechnet hier im Server (Modell bleibt geladen, unter 1 s). Genau
    laeuft als eigener Prozess, damit der Speicher danach wieder frei ist."""
    try:
        import freisteller as F
        a = json.loads(koerper.decode("utf-8-sig"))
        modell = F.name_von(a.get("modell"))
        ordner = P.ordner(name).resolve()
        quelle = (ordner / P.datei_ok(a["datei"])).resolve()   # erst Zeichen pruefen, dann aufloesen
        if ordner not in quelle.parents or not quelle.is_file():
            return {"ok": False, "fehler": "Bild nicht gefunden"}
        if quelle.suffix.lower() == ".svg":
            return {"ok": False, "fehler": "SVG-Grafiken haben keinen Fotohintergrund"}
        if quelle.suffix.lower() not in BILDARTEN:
            return {"ok": False, "fehler": "kein Bild"}
        if not F.pfad(modell):
            return {"ok": False, "fehlt": modell, "mb": F.QUELLEN[modell]["mb"],
                    "fehler": "Freisteller '%s' ist noch nicht eingerichtet" % modell}
        ziel = F.zielname(quelle, modell)
        if ziel.is_file() and not ziel.is_symlink() and ziel.stat().st_mtime >= quelle.stat().st_mtime:
            return {"ok": True, "datei": ziel.relative_to(ordner).as_posix()}   # schon gerechnet
        if modell == "genau":
            if not GENAU_SPERRE.acquire(blocking=False):
                return {"ok": False, "fehler": "es laeuft schon ein gruendliches Freistellen"}
            try:
                r = subprocess.run([sys.executable, str(HIER / "freisteller.py"), str(quelle),
                                    str(ziel), "--modell", "genau"],
                                   capture_output=True, text=True, timeout=1200)
            finally:
                GENAU_SPERRE.release()
            if r.returncode != 0 or not ziel.exists():
                return {"ok": False, "fehler": (r.stderr or "unbekannt").strip().splitlines()[-1][:200]}
        else:
            with SCHNELL_SPERRE:            # unter 1 s, Warten ist unkritisch
                F.freistellen(quelle, ziel, modell)
        return {"ok": True, "datei": ziel.relative_to(ordner).as_posix()}
    except Exception as e:
        return {"ok": False, "fehler": str(e)[:200]}


EIGENE_SEITEN = {"/cockpit.html", "/vorlage.html"}
# Inline-Skripte nur mit ihrer Pruefsumme: eingeschleustes HTML laeuft nicht
SEITEN_CSP = ("default-src 'self'; script-src 'self' %s; style-src 'self' 'unsafe-inline'; "
              "img-src 'self' data: blob:; font-src 'self' data:; connect-src 'self'; "
              "object-src 'none'; base-uri 'none'; form-action 'none'; frame-ancestors 'none'")
_HASHES = {}


def _skript_hashes(weg):
    """'sha256-...' fuer jedes <script> ohne src in der Seite (CRLF wie der Browser zu LF)."""
    datei = HIER / weg.lstrip("/")
    try:
        stand = datei.stat().st_mtime_ns
    except OSError:
        return "'none'"
    if _HASHES.get(weg, (None, ""))[0] != stand:
        s = datei.read_text(encoding="utf-8")
        teile = [x.replace("\r\n", "\n").replace("\r", "\n")
                 for x in re.findall(r"<script>(.*?)</script>", s, re.S)]
        _HASHES[weg] = (stand, " ".join("'sha256-%s'" % base64.b64encode(
            hashlib.sha256(x.encode("utf-8")).digest()).decode() for x in teile) or "'none'")
    return _HASHES[weg][1]


class Kanal(http.server.SimpleHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def end_headers(self):
        self.send_header("Cache-Control", "no-store")
        self.send_header("X-Frame-Options", "DENY")          # nicht in fremde Seiten einrahmbar
        self.send_header("X-Content-Type-Options", "nosniff")
        self.send_header("Cross-Origin-Resource-Policy", "same-origin")   # fremde Seiten betten nichts ein
        # Nur die beiden eigenen Seiten duerfen Skripte ausfuehren, und auch die
        # nur eigene Dateien laden (keine Netzadressen aus einer fremden
        # inhalt.json). Alles andere, auch eine fremde .html in einer Vorlage
        # oder eine hochgeladene SVG, laeuft ohne Skripte.
        weg = urlsplit(self.path).path
        if weg in EIGENE_SEITEN:
            self.send_header("Content-Security-Policy", SEITEN_CSP % _skript_hashes(weg))
        else:
            self.send_header("Content-Security-Policy", "sandbox")
        super().end_headers()

    def translate_path(self, path):
        """URL -> Datei. Code liegt bei HIER, Daten im Karussell-Ordner:
        /projekte/... und /vorlagen/<eigene>/... dort, /fonts/google* aus
        schriften/. Alles andere aus dem Code-Ordner."""
        from urllib.parse import unquote
        teile = [t for t in unquote(urlsplit(path).path).split("/") if t not in ("", ".", "..")]
        if any("\\" in t or ":" in t or P.geraet(t) for t in teile):
            return str(HIER / "__gibt_es_nicht__")
        basis, rest = HIER, teile
        if teile[:1] == ["projekte"]:
            basis, rest = P.PROJEKTE, teile[1:]
        elif teile[:1] == ["vorlagen"] and len(teile) > 1:
            o = P.vorlage_ordner(teile[1])
            basis, rest = (o, teile[2:]) if o else (HIER / "__gibt_es_nicht__", [])
        elif teile[:1] == ["fonts"] and len(teile) > 1 and (teile[1] in ("google.css", "google.json") or teile[1] == "google"):
            basis, rest = P.SCHRIFTEN, teile[1:]
        elif teile[:1] == ["marke"] and len(teile) == 2:          # Logos des Markenpakets
            ok = Path(teile[1]).suffix.lower() in {".png", ".jpg", ".jpeg", ".webp", ".svg"}
            basis, rest = (P.DATEN / "marke", teile[1:]) if ok else (HIER / "__gibt_es_nicht__", [])
        return str(basis.joinpath(*rest))

    def do_HEAD(self):
        if not self._eigen():
            self._text("verboten", 403); return
        super().do_HEAD()

    def _text(self, txt, code=200, art="text/plain"):
        b = txt.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", art + "; charset=utf-8")
        self.send_header("Content-Length", str(len(b)))
        self.end_headers()
        self.wfile.write(b)

    def _weg(self):
        teile = urlsplit(self.path)
        return teile.path.strip("/"), {k: v[0] for k, v in parse_qs(teile.query).items()}

    def _eigen(self):
        """Nur Anfragen vom eigenen Cockpit oder von der Kommandozeile.
        Sperrt fremde Webseiten aus, die im selben Browser offen sind."""
        port = self.server.server_address[1]
        erlaubt = {"127.0.0.1:%d" % port, "localhost:%d" % port}
        if self.headers.get("Host", "") not in erlaubt:
            return False
        if self.headers.get("Sec-Fetch-Site") == "cross-site":   # eingebettet von fremder Seite
            return False
        herkunft = self.headers.get("Origin")
        return herkunft is None or herkunft.replace("http://", "") in erlaubt

    def do_GET(self):
        if not self._eigen():
            self._text("verboten", 403); return
        weg, qs = self._weg()
        if weg == "stand":
            self._text(stand())
            return
        if weg == "projekte":
            self._text(projekte_json(), art="application/json")
            return
        if weg == "vorlagen":
            self._text(vorlagen_json(), art="application/json")
            return
        if weg == "inhalt.json":
            # Alte Tabs und Werkzeuge fragen ohne Projekt: dann das aktuelle.
            try:
                self._text((P.ordner() / "inhalt.json").read_text(encoding="utf-8-sig"),
                           art="application/json")
            except Exception as e:
                self._text("Fehler: %s" % e, 404)
            return
        if weg == "modelle":
            self._text(json.dumps(modelle_stand(), ensure_ascii=False), art="application/json")
            return
        if weg == "marke":
            import marke as MK
            try:
                self._text(json.dumps(MK.lesen()), art="application/json")
            except (ValueError, OSError) as e:
                self._text("Fehler: %s" % " ".join(str(e).split())[:200], 500)
            return
        m = ICON_WEG.fullmatch(weg)
        if m and icon_paket():
            svg = icon_paket().get(m.group(1))
            if svg is None:
                self._text("unbekanntes Icon", 404); return
            self._text(svg, art="image/svg+xml")
            return
        if weg == "karussell.pdf":
            try:
                datei = P.export(qs.get("projekt")) / "karussell.pdf"
                daten = datei.read_bytes()
            except (ValueError, OSError):
                self._text("Fehler: noch kein PDF gebaut", 404); return
            self.send_response(200)
            self.send_header("Content-Type", "application/pdf")
            self.send_header("Content-Disposition",
                             'attachment; filename="%s.pdf"' % datei.parent.name)
            self.send_header("Content-Length", str(len(daten)))
            self.end_headers()
            self.wfile.write(daten)
            return
        if weg == "karussell.zip":
            try:
                ordner = P.export(qs.get("projekt")) / "bilder"
            except ValueError as e:
                self._text("Fehler: %s" % e, 404); return
            puffer = io.BytesIO()
            with zipfile.ZipFile(puffer, "w", zipfile.ZIP_DEFLATED) as z:
                for p in sorted(ordner.glob("slide-*.png")):
                    if re.fullmatch(r"slide-\d+", p.stem):     # keine Finder-/Explorer-Kopien
                        z.write(p, p.name)
            daten = puffer.getvalue()
            self.send_response(200)
            self.send_header("Content-Type", "application/zip")
            self.send_header("Content-Disposition",
                             'attachment; filename="%s.zip"' % ordner.parent.name)
            self.send_header("Content-Length", str(len(daten)))
            self.end_headers()
            self.wfile.write(daten)
            return
        super().do_GET()

    def do_POST(self):
        if not self._eigen():
            self._text("verboten", 403); return
        weg, qs = self._weg()
        n = int(self.headers.get("Content-Length", 0) or 0)
        if n < 0 or n > MAX_UPLOAD:
            self._text("Fehler: zu gross", 413); return
        koerper = self.rfile.read(n)
        try:
            name = qs.get("projekt") or P.aktuelles()
            if weg not in ("projekt", "marke", "markelogo", "schriftdatei"):
                P.ordner(name)          # wirft bei unbekanntem Projekt
        except ValueError as e:
            self._text("Fehler: %s" % e, 404); return
        if weg == "speichern":
            try:
                d = json.loads(koerper.decode("utf-8-sig"))
                assert isinstance(d.get("slides"), list)
            except Exception:
                self._text("Fehler: keine gueltigen Karusselldaten", 400); return
            ziel = P.ordner(name) / "inhalt.json"
            if ziel.is_symlink():
                self._text("Fehler: inhalt.json ist eine Verknuepfung", 400); return
            ziel.write_bytes(koerper)
            self._text(stand_von(name))
        elif weg == "bild":
            roh = qs.get("name") or "bild.png"
            datei = re.sub(r"[^A-Za-z0-9_.\-]", "_", Path(roh).name).strip(". -") or "bild.png"
            if Path(datei).suffix.lower() in WANDELN:
                if sys.platform != "darwin":
                    self._text("Fehler: HEIC- und TIFF-Bilder bitte vorher als JPG oder PNG speichern", 400); return
                try:
                    koerper, datei = _mac_wandeln(koerper, datei)
                except Exception as e:
                    self._text("Fehler: %s" % e, 400); return
            if Path(datei).suffix.lower() not in BILDARTEN:
                self._text("Fehler: nur Bilddateien (%s)" % ", ".join(sorted(BILDARTEN)), 400); return
            try:
                self._text(bild_ablegen(P.ordner(name), datei, koerper))
            except ValueError as e:
                self._text("Fehler: %s" % e, 400)
        elif weg == "ordner":
            try:
                ziel = P.export(name)
                if koerper.decode("utf-8", "ignore").strip() == "pptx":
                    ziel.mkdir(parents=True, exist_ok=True)
                    datei = ziel / "karussell.pptx"
                    oeffnen(ziel, datei if datei.exists() else None)
                elif koerper.decode("utf-8", "ignore").strip() == "pdf":
                    ziel.mkdir(parents=True, exist_ok=True)
                    datei = ziel / "karussell.pdf"
                    oeffnen(ziel, datei if datei.exists() else None)
                else:
                    (ziel / "bilder").mkdir(parents=True, exist_ok=True)
                    oeffnen(ziel / "bilder")
                self._text("geoeffnet")
            except Exception as e:
                self._text("Fehler: %s" % e)
        elif weg == "rendern":
            if not SPERRE.acquire(blocking=False):
                self._text("laeuft schon"); return
            try:
                roh = koerper.decode("utf-8", "ignore").strip() or "bilder"
                # "bilder" = alles, "bilder:3" oder "bilder:1,4" = nur diese Slides
                art, _, wahl = roh.partition(":")
                art = art.strip() or "bilder"
                nur = [int(n) for n in wahl.replace(" ", "").split(",")
                       if n.isdigit()] or None
                self._text(komplett(self.server.server_address[1], art, nur, name))
            except Exception as e:
                self._text("Fehler: %s" % e)
            finally:
                SPERRE.release()
        elif weg == "format":
            if not SPERRE.acquire(blocking=False):
                self._text("laeuft schon"); return
            try:
                self._text(format_umstellen(self.server.server_address[1], name,
                                            koerper.decode("utf-8", "ignore")))
            except Exception as e:
                self._text("Fehler: %s" % e)
            finally:
                SPERRE.release()
        elif weg == "marke":
            # {"aktion": "speichern|logo-weg|logo-ins-projekt|anwenden", ...}
            import marke as MK
            try:
                a = json.loads(koerper.decode("utf-8-sig") or "{}")
                aktion = a.get("aktion")
                if aktion == "speichern":
                    aus = {"ok": True, "marke": MK.schreiben(a.get("marke"))}
                elif aktion == "logo-weg":
                    aus = {"ok": True, "marke": MK.logo_weg(str(a.get("datei", ""))[:100])}
                elif aktion == "logo-ins-projekt":
                    datei, v = MK.logo_ins_projekt(name, a.get("datei"))
                    aus = {"ok": True, "datei": datei, "verhaeltnis": v}
                elif aktion == "anwenden":
                    # nur auf den Stand, den das Cockpit gerade gesichert hat (nichts von k.py ueberschreiben)
                    if isinstance(a.get("stand"), str) and a["stand"] != stand_von(name):
                        raise ValueError("Das Karussell wurde gerade geaendert. Bitte noch einmal anwenden.")
                    meldung = MK.anwenden_auf_projekt(name, farben=a.get("farben") is not False,
                                                      schriften=a.get("schriften") is not False,
                                                      logo=a.get("logo") is True)
                    aus = {"ok": True, "meldung": meldung}
                else:
                    aus = {"ok": False, "fehler": "unbekannte Aktion"}
            except Exception as e:
                aus = {"ok": False, "fehler": " ".join(str(e).split())[:200]}
            self._text(json.dumps(aus), art="application/json")
        elif weg == "markelogo":
            import marke as MK
            try:
                aus = {"ok": True, "marke": MK.logo_ablegen(qs.get("name") or "logo.png", koerper, bild_ablegen)}
            except Exception as e:
                aus = {"ok": False, "fehler": " ".join(str(e).split())[:200]}
            self._text(json.dumps(aus), art="application/json")
        elif weg == "schriftdatei":
            if not SCHRIFT_SPERRE.acquire(blocking=False):
                self._text(json.dumps({"ok": False, "fehler": "es laedt schon eine Schrift"}),
                           art="application/json"); return
            try:
                import schriften
                aus = {"ok": True, "family": schriften.eigene_laden(koerper, qs.get("name") or "schrift.ttf")}
            except Exception as e:
                aus = {"ok": False, "fehler": " ".join(str(e).split())[:200]}
            finally:
                SCHRIFT_SPERRE.release()
            self._text(json.dumps(aus), art="application/json")
        elif weg == "auftrag":
            # Wunsch aus dem Cockpit fuer Claude, gelesen mit `python k.py auftrag`
            try:
                a = json.loads(koerper.decode("utf-8-sig"))
                auftrag = {"zeit": time.strftime("%Y-%m-%d %H:%M"), "folie": int(a["folie"]),
                           "ids": [str(i)[:20] for i in a.get("ids", [])][:60],
                           "text": str(a["text"])[:4000]}
            except Exception:
                self._text("Fehler: ungueltiger Auftrag", 400); return
            ziel = P.auftrag_datei(name)
            ziel.parent.mkdir(parents=True, exist_ok=True)
            ziel.write_text(json.dumps(auftrag, ensure_ascii=False, indent=1), encoding="utf-8")
            self._text("ok")
        elif weg == "bildkopie":
            # Beim Einfuegen aus einem anderen Projekt das Bild mitnehmen
            try:
                a = json.loads(koerper.decode("utf-8-sig"))
                quelle_ordner = P.ordner(a["von"]).resolve()
                datei = P.datei_ok(a["datei"])
                quelle = (quelle_ordner / datei).resolve()
                ziel = (P.ordner(name).resolve() / datei).resolve()
                if quelle_ordner not in quelle.parents or P.ordner(name).resolve() not in ziel.parents:
                    raise ValueError("ausserhalb des Projekts")
                if quelle.suffix.lower() not in BILDARTEN or not quelle.is_file():
                    raise ValueError("kein Bild")
            except Exception:
                self._text("Fehler: Bild nicht gefunden", 400); return
            if not ziel.exists():
                ziel.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(quelle, ziel)
            self._text("ok")
        elif weg == "schrift":
            if not SCHRIFT_SPERRE.acquire(blocking=False):
                self._text(json.dumps({"ok": False, "fehler": "es laedt schon eine Schrift"}),
                           art="application/json"); return
            try:
                import schriften
                fam = json.loads(koerper.decode("utf-8-sig")).get("family", "")
                self._text(json.dumps({"ok": True, "family": schriften.laden(str(fam))}),
                           art="application/json")
            except Exception as e:
                self._text(json.dumps({"ok": False, "fehler": str(e)[:200]}, ensure_ascii=False),
                           art="application/json")
            finally:
                SCHRIFT_SPERRE.release()
        elif weg == "modell":
            try:
                import freisteller as F
                modell = F.name_von(json.loads(koerper.decode("utf-8-sig")).get("modell"))
            except Exception:
                self._text(json.dumps({"ok": False, "fehler": "unbekannter Freisteller"}),
                           art="application/json"); return
            self._text(json.dumps(modell_laden_starten(modell), ensure_ascii=False),
                       art="application/json")
        elif weg == "freistellen":
            self._text(json.dumps(freistellen_auftrag(name, koerper), ensure_ascii=False),
                       art="application/json")
        elif weg == "projekt":
            # {"aktion": "oeffnen|neu|duplizieren", "name": "...", "format": "45|34", "von": "..."}
            try:
                a = json.loads(koerper.decode("utf-8-sig") or "{}")
                aktion = a.get("aktion")
                if aktion == "oeffnen":
                    neu = P.setzen(a.get("name", ""))
                elif aktion == "neu" and a.get("vorlage"):
                    neu = P.neu_aus_vorlage(a.get("name", ""), a["vorlage"])
                    if a.get("marke") is True:
                        import marke as MK
                        try:
                            MK.anwenden_auf_projekt(neu)
                        except Exception:
                            pass                     # dann eben in den Farben der Vorlage
                    neu = P.setzen(neu)
                elif aktion == "neu":
                    neu = P.setzen(P.neu(a.get("name", ""), a.get("format", "45")))
                elif aktion == "vorlage":
                    v = P.als_vorlage(a.get("name", ""), a.get("von") or None)
                    try:
                        vorschau_bauen(self.server.server_address[1], v)
                    except Exception:
                        shutil.rmtree(P.VORLAGEN / v, ignore_errors=True)
                        raise
                    self._text(json.dumps({"ok": True, "vorlage": v}), art="application/json")
                    return
                elif aktion == "duplizieren":
                    neu = P.setzen(P.duplizieren(a.get("name", ""), a.get("von") or None))
                else:
                    raise ValueError("unbekannte Aktion")
                self._text(json.dumps({"ok": True, "name": neu, "stand": stand_von(neu)}),
                           art="application/json")
            except Exception as e:
                self._text(json.dumps({"ok": False, "fehler": str(e)}, ensure_ascii=False),
                           art="application/json")
        else:
            self._text("unbekannt", 404)


class Server(socketserver.ThreadingTCPServer):
    daemon_threads = True
    # Unter Windows erlaubt SO_REUSEADDR einem zweiten Server denselben Port,
    # dann griffe der Ausweich auf einen freien Port nie.
    allow_reuse_address = os.name != "nt"


def server(port=8720):
    """Fester Port, damit der Cockpit-Tab nach einem Neustart weiterlebt."""
    h = functools.partial(Kanal, directory=str(HIER))
    try:
        srv = Server(("127.0.0.1", port), h)
    except OSError:
        srv = Server(("127.0.0.1", 0), h)

    def dauerhaft():
        # Windows kann bei knappem Speicher einen Socketfehler werfen
        # (WinError 10055). Dann kurz warten und weiterlaufen, statt dass
        # das Cockpit still stirbt und Arbeit verloren geht.
        while True:
            try:
                srv.serve_forever()
                return                       # regulaer per shutdown() beendet
            except OSError:
                time.sleep(1)

    threading.Thread(target=dauerhaft, daemon=True).start()
    return srv, srv.server_address[1]


# ---------------------------------------------------------------- Render

# Eine Bildebene freistellen: alles andere unsichtbar, Hintergrund durchsichtig.
# Gibt das sichtbare Rechteck der Ebene auf der Slide zurueck (CSS-Pixel).
_FREISTELLEN = """([id, rand, gast]) => {
  let st = document.getElementById('rasterStil');
  if (!st) {
    st = document.createElement('style'); st.id = 'rasterStil';
    st.textContent = 'html,body,.slide{background:transparent!important}' +
                     '.slide::before,.slide::after{display:none!important}';
    document.head.appendChild(st);
  }
  const slide = document.querySelector('.slide');
  let ziel = null;
  slide.querySelectorAll('.box').forEach(el => {
    const an = el.dataset.art === id && el.classList.contains('gast') === !!gast;
    el.style.visibility = an ? 'visible' : 'hidden';
    if (an) ziel = el;
  });
  if (!ziel) return null;
  const s = slide.getBoundingClientRect(), r = ziel.getBoundingClientRect();
  // rand: Platz fuer einen Schatten, der ueber die Box hinausragt
  const x0 = Math.max(Math.floor(r.left - s.left - rand), 0), y0 = Math.max(Math.floor(r.top - s.top - rand), 0);
  const x1 = Math.min(Math.ceil(r.right - s.left + rand), s.width), y1 = Math.min(Math.ceil(r.bottom - s.top + rand), s.height);
  return (x1 > x0 && y1 > y0) ? {x: x0, y: y0, w: x1 - x0, h: y1 - y0} : null;
}"""

_ZURUECK = """() => {
  const st = document.getElementById('rasterStil'); if (st) st.remove();
  document.querySelectorAll('.slide .box').forEach(el => el.style.visibility = '');
}"""


FILTER_STD = {"hell": 100, "kontrast": 100, "saett": 100, "unschaerfe": 0, "grau": 0, "sepia": 0}
# Wie MASKEN und GERAETE in render.js
MASKEN = {"kreis", "bogen", "herz", "stern", "sechseck", "blob", "raute", "dreieck"}
GERAETE = {"handy", "tablet", "laptop", "browser"}


def zahl(v, std):
    """Zahl aus inhalt.json, bei Unsinn (Text, unendlich) der Standardwert."""
    try:
        f = float(v)
        return f if f == f and abs(f) < 1e9 else std
    except (TypeError, ValueError, OverflowError):
        return std


def _dict(v):
    return v if isinstance(v, dict) else {}


# Was CSS als Hex-Farbe kennt: #rgb, #rgba, #rrggbb, #rrggbbaa
FARBE = re.compile(r"#(?:[0-9a-fA-F]{3,4}|[0-9a-fA-F]{6}|[0-9a-fA-F]{8})")


def farbwert(v, std="#000000"):
    """Gepruefte Farbe als RRGGBB (ohne #) fuer PowerPoint. Auch der
    Rueckfallwert wird geprueft, er kommt oft selbst aus inhalt.json."""
    for x in (v, std, "#000000"):
        if isinstance(x, str) and FARBE.fullmatch(x):
            h = x[1:]
            return ("".join(c * 2 for c in h[:3]) if len(h) in (3, 4) else h[:6]).upper()


def braucht_raster(b):
    """Gleiche Regeln wie bildWirkung() in render.js. SVG kann PIL nicht
    oeffnen, Ausschnitt, Spiegeln, Filter, Deckkraft, Ecken, Rahmen und
    Schatten kann PowerPoint nicht so nachbauen, wie der Browser sie zeigt."""
    if b.get("typ") == "form":
        return True                     # Formen und Icons: der Browser zeichnet sie
    if b.get("typ") not in ("bild", "form"):
        # Text mit Bild in der Schrift kann PowerPoint nicht. Nahtloser Text wird ein
        # Bild, damit beide Haelften an der Folienkante exakt zusammenpassen (echter
        # Text saesse in Canva ein paar Pixel anders als die gerasterte Haelfte).
        # Der Rest bleibt echter, bearbeitbarer Text.
        bf = b.get("bildfuellung")
        frei = b.get("x") is not None and b.get("y") is not None
        return (isinstance(bf, str) and bool(P.DATEI.fullmatch(bf))) or (bool(b.get("nahtlos")) and frei)
    if isinstance(b.get("geraet"), str) and b["geraet"] in GERAETE:
        return True                     # Geraete-Rahmen, auch leer (Bildschirm bleibt schwarz)
    if not b.get("datei"):
        return False                    # leerer Rahmen: auf der Folie nicht sichtbar
    if isinstance(b.get("maske"), str) and b["maske"] in MASKEN:
        return True
    if str(b.get("datei", "")).lower().endswith(".svg"):
        return True
    a = _dict(b.get("ausschnitt"))
    if zahl(a.get("z"), 1) != 1 or zahl(a.get("x"), 50) != 50 or zahl(a.get("y"), 50) != 50:
        return True
    f = {**FILTER_STD, **_dict(b.get("filter"))}
    if any(zahl(f[k], v) != v for k, v in FILTER_STD.items()):
        return True
    # Schatten ohne eigene Werte ist {} und trotzdem an
    if b.get("spiegeln") or b.get("ecken") or b.get("schatten") is not None:
        return True
    if b.get("deckkraft") is not None and zahl(b["deckkraft"], 100) < 100:
        return True
    return bool(zahl(_dict(b.get("rahmen")).get("breite"), 0))


def _schattenrand(b):
    if b.get("schatten") is None:
        return 0
    s = _dict(b.get("schatten"))
    return int(min(400, abs(zahl(s.get("abstand"), 14)) + 3 * abs(zahl(s.get("weich"), 24)))) + 2


def _gaeste_pruefen(roh, anzahl):
    """window.GAESTE aus vorlage.html: nur saubere Eintraege, hoechstens 200."""
    aus, gesehen = [], set()
    for g in (roh if isinstance(roh, list) else [])[:200]:
        if not isinstance(g, dict):
            continue
        art, bid, folie, lage = g.get("art"), g.get("id"), g.get("folie"), g.get("lage")
        if isinstance(art, str) and isinstance(bid, str) and type(folie) is int and \
                1 <= folie <= anzahl and lage in ("unten", "oben") and len(art) <= 300 and \
                len(bid) <= 200 and art not in gesehen:
            gesehen.add(art)
            aus.append({"art": art, "id": bid, "folie": folie, "lage": lage})
    return aus


def _rastern(pg, slide, nr, ordner, gaeste=(), alle=()):
    """Der Browser zeichnet jede Bildebene, die PowerPoint nicht nachbauen
    kann, so wie sie auf der Slide steht, als durchsichtiges PNG (siehe
    braucht_raster). Zuschnitt, Effekte und Drehung sind darin schon erledigt."""
    aus = {}
    for j, b in enumerate(slide.get("bloecke", [])[:200]):     # wie MAX_BLOECKE in render.js
        if not braucht_raster(b):
            continue
        bid = b.get("id") or "b%d" % (j + 1)      # wie renderSlide in render.js
        r = pg.evaluate(_FREISTELLEN, [bid, _schattenrand(b), False])
        if not r:
            continue
        pfad = ordner / f"{nr:02d}-{j + 1:02d}.png"   # nie die ID aus inhalt.json im Pfad
        pg.screenshot(path=str(pfad), omit_background=True,
                      clip={"x": r["x"], "y": r["y"], "width": r["w"], "height": r["h"]})
        aus[bid] = {**r, "pfad": str(pfad)}
    # Nahtlose Bloecke anderer Folien, soweit sie hier hineinragen: als Bild
    for j, g in enumerate(gaeste):
        heim = {}
        if 1 <= g["folie"] <= len(alle):
            for jj, b in enumerate(alle[g["folie"] - 1].get("bloecke", [])):
                if str(b.get("id") or "b%d" % (jj + 1)) == g["id"]:
                    heim = b
                    break
        r = pg.evaluate(_FREISTELLEN, [g["art"], _schattenrand(heim), True])
        if not r:
            continue
        pfad = ordner / f"{nr:02d}-g{j + 1:02d}.png"
        pg.screenshot(path=str(pfad), omit_background=True,
                      clip={"x": r["x"], "y": r["y"], "width": r["w"], "height": r["h"]})
        aus[("gast", g["art"])] = {**r, "pfad": str(pfad)}   # Tupel: kollidiert nie mit einer Block-ID
    if aus:
        pg.evaluate(_ZURUECK)
    return aus


def messen(slides, stil, port, bilder=True, nur=None, name=None, raster=False):
    """Jede Slide im Browser aufbauen, Masse zurueckgeben, auf Wunsch als PNG sichern.

    nur: Liste von Slide-Nummern (1-basiert). Dann werden ausschliesslich diese
    gerendert und die uebrigen PNG bleiben unangetastet. Spart Zeit und Rechnung,
    wenn nur eine Slide geaendert wurde.
    raster: Bildebenen mit SVG oder Effekten fuer die Canva-Datei als PNG in doppelter Aufloesung
    nach projekte/<name>/.raster/ zeichnen (siehe _rastern)."""
    name = name or P.aktuelles()
    aus = P.export(name)
    ziel = aus / "bilder"
    (ziel if bilder else aus).mkdir(parents=True, exist_ok=True)
    rordner = P.ordner(name) / ".raster"
    if raster:
        if os.path.lexists(rordner) and P._verknuepft(rordner):
            raise ValueError("projekte/%s/.raster ist eine Verknuepfung, Export abgebrochen" % P.ordner(name).name)
        rordner.mkdir(exist_ok=True)
        for alt in rordner.glob("*.png"):       # eigener Zwischenspeicher, immer frisch
            alt.unlink()
    masse = []
    B, H = format_ok(stil)
    groesse = {"width": B, "height": H}
    with sync_playwright() as p:
        b = _browser(p)
        pg = b.new_page(viewport=groesse, device_scale_factor=1) if bilder else None
        rp = b.new_page(viewport=groesse, device_scale_factor=2) if raster else None
        auswahl = range(len(slides)) if not nur else [n - 1 for n in nur
                                                      if 1 <= n <= len(slides)]
        for i in auswahl:
            url = f"http://127.0.0.1:{port}/vorlage.html?projekt={name}&render={i}"
            for seite in (pg, rp):
                if seite:
                    seite.goto(url)
                    seite.wait_for_function(FERTIG, timeout=20000)
            m = (pg or rp).evaluate("() => window.MASSE")
            m["gaeste"] = _gaeste_pruefen((pg or rp).evaluate("() => (window.GAESTE || []).slice(0, 200)"), len(slides))
            if bilder:
                pg.screenshot(path=str(ziel / f"slide-{i+1:02d}.png"))
            if raster:
                m["raster"] = _rastern(rp, slides[i], i + 1, rordner, m["gaeste"], slides)
            masse.append(m)
        b.close()
    if bilder and not nur:
        for a in sorted(ziel.glob("slide-*.png")):
            m = re.fullmatch(r"slide-(\d+)", a.stem)          # "slide-01 Kopie.png" nicht anfassen
            if m and int(m.group(1)) > len(slides):
                a.unlink()
    return masse


def _hl(run, rgb):
    rPr = run._r.get_or_add_rPr()
    hl  = rPr.makeelement(qn("a:highlight"), {})
    clr = rPr.makeelement(qn("a:srgbClr"), {"val": rgb})
    hl.append(clr)
    try:
        rPr.insert_element_before(hl, "a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill",
                                  "a:latin", "a:ea", "a:cs", "a:sym",
                                  "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")
    except Exception:
        rPr.append(hl)


def _run(p, txt, art, groesse, stil):
    """art ist ein Formatsatz: b/i/u = wahr, m = markiert, c = Farbe, f = Schrift."""
    art = art or {}
    r = p.add_run(); r.text = txt
    f = r.font
    f.size = Pt(max(1.0, min(zahl(groesse, 40) * 0.75, 4000.0)))   # python-pptx erlaubt 1 bis 4000 pt
    f.name = art.get("f") or stil.get("schrift") or "Montserrat"
    f.bold = bool(art.get("b"))
    f.italic = bool(art.get("i"))
    f.underline = bool(art.get("u"))
    c = art.get("c")
    if c == "akzent":
        c = stil["akzent"]
    if art.get("m"):
        _hl(r, farbwert(stil["akzent"], "#718d81"))
        f.color.rgb = RGBColor.from_string(farbwert(c, "#ffffff")) if c else RGBColor(0xFF, 0xFF, 0xFF)
    else:
        f.color.rgb = RGBColor.from_string(farbwert(c, stil["text"]) if c else farbwert(stil["text"], "#313538"))
    _texteffekte(r, groesse, stil)


_NACH_EFFEKT = ("a:highlight", "a:uLnTx", "a:uLn", "a:uFillTx", "a:uFill", "a:latin", "a:ea",
                "a:cs", "a:sym", "a:hlinkClick", "a:hlinkMouseOver", "a:rtl", "a:extLst")


def _texteffekte(r, groesse, stil):
    """Text-Effekte als echte PowerPoint-Eigenschaften, damit der Text in Canva
    bearbeitbar bleibt. Gegenstueck zu textWirkung() in render.js."""
    rPr = r._r.get_or_add_rPr()
    z = max(-1000, min(zahl(stil.get("zeichen"), 0), 5000))
    if z:                                           # em/1000 -> hundertstel Punkt, Schema: +-400000
        rPr.set("spc", str(max(-400000, min(int(round(z / 1000 * groesse * 0.75 * 100)), 400000))))
    if stil.get("gross"):
        rPr.set("cap", "all")
    u = _dict(stil.get("umriss"))
    ub = max(0, min(zahl(u.get("breite"), 0), 200))
    if ub:
        hohl = bool(u.get("hohl"))
        ln = rPr.makeelement(qn("a:ln"), {"w": str(int((ub if hohl else ub * 2) * PX))})
        sf = ln.makeelement(qn("a:solidFill"), {})
        sf.append(sf.makeelement(qn("a:srgbClr"), {"val": farbwert(u.get("farbe"), "#000000")}))
        ln.append(sf)
        rPr.insert(0, ln)
        if hohl:                                    # nur der Umriss, innen durchsichtig
            alt = rPr.find(qn("a:solidFill"))
            if alt is not None:
                i = list(rPr).index(alt)
                rPr.remove(alt)
                rPr.insert(i, rPr.makeelement(qn("a:noFill"), {}))
    if stil.get("schatten") is not None:
        sch = _dict(stil.get("schatten"))
        eff = rPr.makeelement(qn("a:effectLst"), {})
        sh = eff.makeelement(qn("a:outerShdw"), {
            "blurRad": str(int(min(abs(zahl(sch.get("weich"), 10)), 400) * PX)),
            "dist": str(int(min(abs(zahl(sch.get("abstand"), 6)), 400) * PX)),
            "dir": str(int((zahl(sch.get("winkel"), 45) % 360) * 60000)),
            "algn": "ctr", "rotWithShape": "0"})
        clr = sh.makeelement(qn("a:srgbClr"), {"val": farbwert(sch.get("farbe"), "#000000")})
        clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(max(0, min(100, zahl(sch.get("deck"), 50))) * 1000))}))
        sh.append(clr)
        eff.append(sh)
        rPr.insert_element_before(eff, *_NACH_EFFEKT)


def _einzug(p, px):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(px * PX)))
    pPr.set("indent", str(-int(px * PX)))
    pPr.append(pPr.makeelement(qn("a:buNone"), {}))


AUSR = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}


def _textbox(sl, p, H):
    dreh = zahl(p.get("dreh"), 0) % 360
    x, y, w = (int(max(-20000, min(zahl(p.get(k), 0), 20000))) for k in ("x", "y", "w"))
    # Ungedreht darf der Kasten bis zum Folienrand reichen (Luft fuer andere
    # Schriftmasse in Canva). Gedreht muss er genau so hoch sein wie im
    # Browser, sonst liegt der Drehpunkt woanders.
    flaeche = p.get("flaeche")
    hoehe = zahl(p.get("h"), 0) or (zahl(p.get("hBox"), 0) if (dreh or flaeche is not None) else 0) or (H - y)
    hoehe = min(hoehe, 20000)
    tb = sl.shapes.add_textbox(Emu(x * PX), Emu(y * PX), Emu(max(1, w) * PX), Emu(int(max(40, hoehe)) * PX))
    if dreh:
        tb.rotation = dreh
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    if flaeche is not None:                       # Hintergrundflaeche hinter dem Text
        fl = _dict(flaeche)
        tb.fill.solid()
        tb.fill.fore_color.rgb = RGBColor.from_string(farbwert(fl.get("farbe"), "#ffffff"))
        deck = zahl(fl.get("deck"), 100)
        if deck < 100:
            clr = tb.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            clr.append(clr.makeelement(qn("a:alpha"), {"val": str(int(max(0, deck) * 1000))}))
        innen = Emu(int(max(0, min(zahl(fl.get("innen"), 24), 500)) * PX))
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = innen
        rund = zahl(fl.get("rund"), 0)
        if rund > 0:
            geo = tb._element.spPr.find(qn("a:prstGeom"))
            geo.set("prst", "roundRect")
            av = geo.find(qn("a:avLst"))
            if av is None:
                av = geo.makeelement(qn("a:avLst"), {}); geo.append(av)
            kurz = max(1, min(max(1, w), int(max(40, hoehe))))
            av.append(av.makeelement(qn("a:gd"), {"name": "adj",
                                                   "fmla": "val %d" % min(50000, rund / kurz * 100000)}))
    return tf


def _bild(sl, pfad, p):
    """Bild einsetzen und wie CSS background-size:cover beschneiden."""
    iw, ih = Image.open(pfad).size
    w = max(1, int(min(zahl(p.get("w"), 1), 20000)))
    h = max(1, int(min(zahl(p.get("h"), 0), 20000)) or round(w * ih / max(1, iw)))
    x, y = (int(max(-20000, min(zahl(p.get(k), 0), 20000))) for k in ("x", "y"))
    pic = sl.shapes.add_picture(str(pfad), Emu(x * PX), Emu(y * PX), Emu(w * PX), Emu(h * PX))
    if zahl(p.get("dreh"), 0):
        pic.rotation = zahl(p.get("dreh"), 0) % 360
    ziel, quelle = w / h, iw / ih
    if quelle > ziel:
        f = (1 - ziel / quelle) / 2
        pic.crop_left = pic.crop_right = f
    elif quelle < ziel:
        f = (1 - quelle / ziel) / 2
        pic.crop_top = pic.crop_bottom = f


def _raster(sl, r):
    """Vom Browser gezeichnete Ebene (SVG) genau dorthin, wo sie stand. Kein
    Zuschnitt mehr, der ist im Bild schon enthalten."""
    sl.shapes.add_picture(r["pfad"], Emu(int(r["x"]) * PX), Emu(int(r["y"]) * PX),
                          Emu(int(r["w"]) * PX), Emu(int(r["h"]) * PX))


def _mitFont(stil, p):
    """Elementeigene Schrift und Farbe durchreichen, sonst die globalen."""
    s2 = None
    if p.get("font"):
        s2 = dict(stil); s2["schrift"] = p["font"]
    if p.get("farbe"):
        s2 = dict(s2 or stil); s2["text"] = p["farbe"]
    for k in ("zeichen", "gross", "schatten", "umriss"):
        if p.get(k) is not None:
            s2 = dict(s2 or stil); s2[k] = p[k]
    return s2 or stil


def _ebenen(s):
    """Ebenenfolge einer Slide von unten nach oben, wie im Cockpit."""
    bl = s.get("bloecke", [])
    da = [b.get("id") for b in bl]
    folge = [k for k in (s.get("ebenen") or []) if k in da]
    for b in bl:
        if b.get("typ") == "bild" and b.get("x") is None and b.get("id") not in folge:
            folge.append(b["id"])
    for k in da:
        if k not in folge:
            folge.append(k)
    return folge


def _fliesstext(tf, text, groesse, zeilenabstand, ausr, stil):
    """Absaetze, Aufzaehlungen und Auszeichnungen in einen Textrahmen schreiben."""
    erste = True
    for i, block in enumerate(re.split(r"\n\s*\n", text)):
        zeilen = block.split("\n")
        liste = all(z.strip().startswith("- ") for z in zeilen)
        for j, zeile in enumerate(zeilen if liste else [block]):
            p = tf.paragraphs[0] if erste else tf.add_paragraph()
            erste = False
            p.line_spacing = zeilenabstand
            p.alignment = ausr
            if j == 0 and i > 0:
                p.space_before = Pt(groesse * 0.85 * 0.75)
            if liste:
                _einzug(p, groesse * 1.35)
                _run(p, "•  ", {}, groesse, stil)
                zeile = zeile.strip()[2:]
            for kk, teil in enumerate(zeile.split("\n")):
                if kk:
                    p.add_line_break()
                for txt, art in segmente(teil):
                    _run(p, txt, art, groesse, stil)


def _hintergrund(sl, hg, stil, quelle, B, H):
    """Folienhintergrund nativ: Farbe, echter Verlauf oder Bild ganz unten."""
    fl = sl.background.fill
    grund = farbwert(hg.get("farbe"), stil["bg"])
    if hg.get("art") == "verlauf":
        fl.gradient()
        st = fl.gradient_stops
        st[0].color.rgb = RGBColor.from_string(grund); st[0].position = 0.0
        st[1].color.rgb = RGBColor.from_string(farbwert(hg.get("farbe2"), "#000000")); st[1].position = 1.0
        # CSS: 0 Grad = nach oben, im Uhrzeigersinn. PowerPoint: 0 = nach rechts.
        uhr = (zahl(hg.get("winkel"), 180) - 90) % 360
        fl.gradient_angle = (360 - uhr) % 360       # python-pptx zaehlt gegen den Uhrzeigersinn
        return
    fl.solid()
    fl.fore_color.rgb = RGBColor.from_string(grund)
    d = hg.get("datei")
    p = quelle / d if isinstance(d, str) and P.DATEI.fullmatch(d) else None
    # Nur echte Rasterbilder (SVG kann PIL nicht oeffnen, NUL "existiert" unter Windows)
    if hg.get("art") == "bild" and p and p.is_file() and p.suffix.lower() in FORMAT:
        try:
            _bild(sl, p, {"x": 0, "y": 0, "w": B, "h": H})
        except Exception:
            pass                                    # die Farbe darunter bleibt stehen


def pptx_bauen(slides, stil, masse, name=None):
    """Jeder Block wird ein eigenes, frei positioniertes PPTX-Objekt."""
    quelle = P.ordner(name)
    B, H = stil["breite"], stil["hoehe"]
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(B * PX), Emu(H * PX)
    leer = prs.slide_layouts[6]

    for s, m in zip(slides, masse):
        sl = prs.slides.add_slide(leer)
        _hintergrund(sl, _dict(s.get("hg")), stil, quelle, B, H)
        L = m.get("layout") or {}
        karte = {b.get("id"): b for b in s.get("bloecke", [])}
        gaeste, rastern = m.get("gaeste") or [], m.get("raster") or {}
        for g in gaeste:                      # nahtlos von frueheren Folien: ganz unten
            if g.get("lage") == "unten" and rastern.get(("gast", g.get("art"))):
                _raster(sl, rastern[("gast", g["art"])])

        # Einfuegereihenfolge ist die Stapelung: zuerst = ganz hinten.
        for k in _ebenen(s):
            b, p = karte.get(k), L.get(k)
            if not b or not p:
                continue
            r = (m.get("raster") or {}).get(k)
            if r:                                 # vom Browser gezeichnet, gleich welcher Art
                _raster(sl, r)
                continue
            if b.get("typ") in ("bild", "form"):
                if braucht_raster(b):
                    pass                      # nicht sichtbar auf der Slide
                elif b.get("datei") and P.DATEI.fullmatch(str(b["datei"])) and \
                        (quelle / b["datei"]).is_file() and (quelle / b["datei"]).suffix.lower() in FORMAT:
                    _bild(sl, quelle / b["datei"], p)
                continue
            if not b.get("text") or (b.get("nahtlos") and b.get("x") is not None):
                continue                      # nahtloser Text ohne Rasterbild liegt ganz neben der Folie
            tf = _textbox(sl, p, H)
            _fliesstext(tf, b["text"], p.get("groesse", stil["textgroesse"]),
                        p.get("abstand", stil["textAbstand"]),
                        AUSR.get(p.get("aus", "left"), PP_ALIGN.LEFT),
                        _mitFont(stil, p))
            if p.get("fett"):
                for pa in tf.paragraphs:
                    for r in pa.runs:
                        r.font.bold = True
        for g in gaeste:                      # nahtlos von spaeteren Folien: oben
            if g.get("lage") == "oben" and rastern.get(("gast", g.get("art"))):
                _raster(sl, rastern[("gast", g["art"])])

    prs.save(str(P.export(name) / "karussell.pptx"))


def kontaktbogen(n, stil, name=None, spalten=4):
    aus = P.export(name)
    kb = 270
    kh = round(kb * stil["hoehe"] / stil["breite"])
    luft, fuss = 14, 26
    zeilen = (n + spalten - 1) // spalten
    bild = Image.new("RGB", (spalten * kb + (spalten + 1) * luft,
                             zeilen * (kh + fuss) + (zeilen + 1) * luft), "#5f6668")
    d = ImageDraw.Draw(bild)
    try:
        fnt = ImageFont.truetype(str(HIER / "fonts" / "Montserrat.ttf"), 16)
    except Exception:
        fnt = ImageFont.load_default()
    for i in range(n):
        x = luft + (i % spalten) * (kb + luft)
        y = luft + (i // spalten) * (kh + fuss + luft)
        s = Image.open(aus / "bilder" / f"slide-{i+1:02d}.png").resize((kb, kh), Image.LANCZOS)
        bild.paste(s, (x, y))
        d.text((x + 2, y + kh + 4), f"{i+1} / {n}", font=fnt, fill="#ffffff")
    bild.save(aus / "kontaktbogen.png")


def komplett(port, art="alles", nur=None, name=None):
    """art: 'bilder' fuer Instagram, 'pptx' fuer Canva, 'alles' fuer die Kommandozeile."""
    name = name or P.aktuelles()
    if art == "pdf":
        return pdf_bauen(port, name)
    slides, stil = laden(name)
    canva = art in ("pptx", "alles")
    masse = messen(slides, stil, port, bilder=(art != "pptx"), nur=None if canva else nur,
                   name=name, raster=canva)
    if art in ("pptx", "alles"):
        pptx_bauen(slides, stil, masse, name)
    if art == "alles":
        kontaktbogen(len(slides), stil, name)
    if art == "pptx":
        return "Canva-Datei fertig (karussell.pptx)"
    if nur:
        return f"Slide {', '.join(str(n) for n in nur)} neu ({stil['breite']}x{stil['hoehe']})"
    return f"{len(slides)} Bilder fertig ({stil['breite']}x{stil['hoehe']})"


def belegt(port):
    """Laeuft dort schon ein Cockpit?"""
    import socket
    with socket.socket() as s:
        s.settimeout(0.3)
        return s.connect_ex(("127.0.0.1", port)) == 0


if __name__ == "__main__":
    # Zweiter Start soll kein zweites Cockpit aufmachen, sondern das offene zeigen.
    if "--cockpit" in sys.argv and belegt(8720):
        print("Cockpit laeuft bereits: http://127.0.0.1:8720/cockpit.html")
        browser_auf("http://127.0.0.1:8720/cockpit.html")
        sys.exit()

    if "--cockpit" in sys.argv:
        if P.einstellungen_kaputt():
            print("Die Kit-Einstellungen (%s) sind beschaedigt. Sag Claude: "
                  "reparier die Kit-Einstellungen. Das Cockpit startet erst danach," % P.EINST)
            print("damit deine Karussells nicht in einem anderen Ordner gesucht werden.")
            sys.exit(1)
        if not P.gesetzt():
            print("Karussell-Ordner:", P.ordner_setzen(P.vorschlag()))
        try:
            import einrichten
            einrichten.startdatei()
        except Exception as e:
            print("Hinweis: Startdatei nicht erneuert (%s)" % e)
    srv, port = server()
    if "--pdf" in sys.argv:
        print(pdf_bauen(port, P.aus_argv(sys.argv)))
        srv.shutdown()
        sys.exit()
    if "--cockpit" in sys.argv or "--vorschau" in sys.argv:
        seite = "cockpit.html" if "--cockpit" in sys.argv else "vorlage.html"
        url = f"http://127.0.0.1:{port}/{seite}"
        print(url)
        print("Laeuft. Fenster offen lassen, %s beendet." % ("ctrl+C" if sys.platform == "darwin" else "Strg+C"))
        browser_auf(url)
        try:
            threading.Event().wait()
        except KeyboardInterrupt:
            srv.shutdown()
        sys.exit()
    print(komplett(port, name=P.aus_argv(sys.argv)))
    srv.shutdown()
